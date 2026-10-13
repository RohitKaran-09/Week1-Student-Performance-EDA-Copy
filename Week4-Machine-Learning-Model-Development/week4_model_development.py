"""Week 4: Machine Learning Model Development and Evaluation.

Task: predict whether a student reaches the descriptive G3 >= 10 threshold from
non-grade information in the official UCI Student Performance Mathematics data.

Leakage prevention: G1 and G2 are intentionally excluded from predictors.
No synthetic fallback data is used.
"""
from io import BytesIO
from pathlib import Path
import json
import urllib.request
import zipfile

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns

from sklearn.compose import ColumnTransformer
from sklearn.dummy import DummyClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.impute import SimpleImputer
from sklearn.inspection import permutation_importance
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score, balanced_accuracy_score, precision_score, recall_score,
    f1_score, roc_auc_score, confusion_matrix, classification_report,
    RocCurveDisplay
)
from sklearn.model_selection import (
    train_test_split, StratifiedKFold, cross_validate
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.tree import DecisionTreeClassifier

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
OUT_DIR = ROOT / "outputs"
PLOT_DIR = ROOT / "visualizations"
REPORT_PATH = ROOT / "Week4_Machine_Learning_Model_Development_Report.docx"
SUBMISSION_DOCX = ROOT / "Submission_Description.docx"
SUBMISSION_TXT = ROOT / "submission_description.txt"

UCI_ZIP = "https://archive.ics.uci.edu/static/public/320/student+performance.zip"
UCI_PAGE = "https://archive.ics.uci.edu/dataset/320/student%2Bperformance"
DOI = "https://doi.org/10.24432/C5TG7T"
RANDOM_STATE = 42
TEST_SIZE = 0.20
PASS_MARK = 10

def find_math_csv(zf):
    for name in zf.namelist():
        if name.lower().endswith("student-mat.csv"):
            return zf.read(name)
    for name in zf.namelist():
        if name.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(BytesIO(zf.read(name))) as nested:
                    found = find_math_csv(nested)
                    if found is not None:
                        return found
            except zipfile.BadZipFile:
                continue
    return None

def acquire_data():
    candidates = [
        ROOT.parent / "data" / "student-mat.csv",
        ROOT / "data" / "student-mat.csv",
        ROOT.parent.parent / "data" / "student-mat.csv",
    ]
    for path in candidates:
        if path.exists():
            frame = pd.read_csv(path, sep=";")
            if len(frame) and {"G3", "G1", "G2", "studytime", "failures"}.issubset(frame.columns):
                return frame, "Existing official UCI Mathematics CSV from the Week 1 project"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(UCI_ZIP, headers={"User-Agent": "Week4-Student-ML/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as response:
            archive_bytes = response.read()
        with zipfile.ZipFile(BytesIO(archive_bytes)) as archive:
            payload = find_math_csv(archive)
        if payload is None:
            raise RuntimeError("student-mat.csv not found in UCI archive")
        (DATA_DIR / "student-mat.csv").write_bytes(payload)
        frame = pd.read_csv(BytesIO(payload), sep=";")
        if frame.empty or not {"G3", "G1", "G2", "studytime", "failures"}.issubset(frame.columns):
            raise RuntimeError("Downloaded data failed schema validation")
        return frame, "Downloaded official UCI archive: " + UCI_ZIP
    except Exception as exc:
        raise RuntimeError("Official UCI data could not be acquired; no synthetic fallback was used.") from exc

def prepare_data(raw):
    df = raw.copy()
    initial_rows = len(df)
    duplicates = int(df.duplicated().sum())
    df = df.drop_duplicates().copy()
    for col in df.select_dtypes(include=["object", "string"]).columns:
        df[col] = df[col].astype("string").str.strip()
    for col in ["age", "studytime", "failures", "absences", "G1", "G2", "G3"]:
        if col in df:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    if "G3" not in df:
        raise ValueError("The source is missing the final grade G3.")
    missing_target = int(df["G3"].isna().sum())
    rows_before_target_drop = len(df)
    df = df.dropna(subset=["G3"]).copy()
    rows_removed_missing_target = rows_before_target_drop - len(df)
    invalid_grade = ~df["G3"].between(0, 20)
    invalid_grade_count = int(invalid_grade.sum())
    df = df.loc[~invalid_grade].copy()
    df["target_pass"] = (df["G3"] >= PASS_MARK).astype(int)
    if df["target_pass"].nunique() < 2:
        raise ValueError("The target contains only one class after preparation.")
    audit = {
        "source_rows": int(initial_rows),
        "duplicate_rows_removed": duplicates,
        "missing_final_grades_before_filter": missing_target,
        "rows_removed_for_missing_final_grade": int(rows_removed_missing_target),
        "invalid_final_grade_values_removed": invalid_grade_count,
        "rows_after_preparation": int(len(df)),
        "target_definition": f"1 if G3 >= {PASS_MARK}; otherwise 0",
        "predictor_exclusions": ["G1", "G2", "G3", "target_pass"],
        "synthetic_fallback_used": False,
        "note": "Rows with a missing target or invalid final grade are excluded. Predictor imputation is fitted inside training-fold pipelines."
    }
    return df, audit

def build_preprocessor(X):
    numeric_cols = X.select_dtypes(include=["number", "bool"]).columns.tolist()
    categorical_cols = [c for c in X.columns if c not in numeric_cols]
    transformers = []
    if numeric_cols:
        numeric_pipeline = Pipeline(steps=[
            ("imputer", SimpleImputer(strategy="median")),
            ("scaler", StandardScaler())
        ])
        transformers.append(("numeric", numeric_pipeline, numeric_cols))
    if categorical_cols:
        categorical_pipeline = Pipeline(steps=[
            ("imputer", SimpleImputer(strategy="most_frequent")),
            ("onehot", OneHotEncoder(handle_unknown="ignore"))
        ])
        transformers.append(("categorical", categorical_pipeline, categorical_cols))
    return ColumnTransformer(transformers=transformers, remainder="drop"), numeric_cols, categorical_cols

def build_models(preprocessor):
    return {
        "Dummy baseline": Pipeline([
            ("preprocess", preprocessor),
            ("model", DummyClassifier(strategy="prior"))
        ]),
        "Logistic regression": Pipeline([
            ("preprocess", preprocessor),
            ("model", LogisticRegression(max_iter=3000, class_weight="balanced", random_state=RANDOM_STATE))
        ]),
        "Decision tree": Pipeline([
            ("preprocess", preprocessor),
            ("model", DecisionTreeClassifier(max_depth=4, min_samples_leaf=8, class_weight="balanced", random_state=RANDOM_STATE))
        ]),
        "Random forest": Pipeline([
            ("preprocess", preprocessor),
            ("model", RandomForestClassifier(n_estimators=250, min_samples_leaf=4, class_weight="balanced", random_state=RANDOM_STATE, n_jobs=1))
        ]),
    }

def evaluate_models(X_train, y_train, X_test, y_test):
    preprocessor, numeric_cols, categorical_cols = build_preprocessor(X_train)
    models = build_models(preprocessor)
    cv = StratifiedKFold(n_splits=5, shuffle=True, random_state=RANDOM_STATE)
    cv_rows = []
    fit_models = {}
    for name, pipeline in models.items():
        scores = cross_validate(
            pipeline, X_train, y_train, cv=cv,
            scoring={"roc_auc": "roc_auc", "f1": "f1", "balanced_accuracy": "balanced_accuracy"},
            return_train_score=False, n_jobs=1, error_score="raise"
        )
        cv_rows.append({
            "model": name,
            "cv_roc_auc_mean": float(np.mean(scores["test_roc_auc"])),
            "cv_roc_auc_std": float(np.std(scores["test_roc_auc"], ddof=1)),
            "cv_f1_mean": float(np.mean(scores["test_f1"])),
            "cv_balanced_accuracy_mean": float(np.mean(scores["test_balanced_accuracy"])),
        })
        fit_models[name] = pipeline
    cv_table = pd.DataFrame(cv_rows).sort_values("cv_roc_auc_mean", ascending=False).reset_index(drop=True)
    selected_name = str(cv_table.iloc[0]["model"])
    selected = fit_models[selected_name]
    selected.fit(X_train, y_train)
    pred = selected.predict(X_test)
    prob = selected.predict_proba(X_test)[:, 1]
    metrics = {
        "selected_model": selected_name,
        "selection_metric": "Mean 5-fold stratified training cross-validation ROC-AUC",
        "test_size_fraction": TEST_SIZE,
        "random_state": RANDOM_STATE,
        "test_count": int(len(y_test)),
        "train_count": int(len(y_train)),
        "positive_class_name": f"G3 >= {PASS_MARK}",
        "negative_class_name": f"G3 < {PASS_MARK}",
        "positive_class_rate_all_pct": float(100 * (y_train.sum() + y_test.sum()) / (len(y_train) + len(y_test))),
        "test_accuracy": float(accuracy_score(y_test, pred)),
        "test_balanced_accuracy": float(balanced_accuracy_score(y_test, pred)),
        "test_precision": float(precision_score(y_test, pred, zero_division=0)),
        "test_recall": float(recall_score(y_test, pred, zero_division=0)),
        "test_f1": float(f1_score(y_test, pred, zero_division=0)),
        "test_roc_auc": float(roc_auc_score(y_test, prob)),
        "confusion_matrix_labels": ["G3 < 10", "G3 >= 10"],
        "confusion_matrix": confusion_matrix(y_test, pred, labels=[0, 1]).tolist(),
        "classification_report": classification_report(y_test, pred, labels=[0, 1], target_names=[f"Below {PASS_MARK}", f"At or above {PASS_MARK}"], output_dict=True, zero_division=0),
        "class_balance_all": {str(int(k)): int(v) for k, v in pd.Series(np.concatenate([y_train.to_numpy(), y_test.to_numpy()])).value_counts().sort_index().items()},
        "cv_results": cv_table.to_dict(orient="records"),
        "numeric_predictors": numeric_cols,
        "categorical_predictors": categorical_cols,
    }
    return selected, pred, prob, cv_table, metrics

def make_plots(y_test, pred, prob, cv_table, model, X_test, y_train, metrics):
    PLOT_DIR.mkdir(parents=True, exist_ok=True)
    sns.set_theme(style="whitegrid", context="notebook")
    paths = []

    # Model comparison by CV ROC AUC.
    fig, ax = plt.subplots(figsize=(9, 5.2))
    sorted_table = cv_table.sort_values("cv_roc_auc_mean", ascending=True)
    ax.barh(sorted_table["model"], sorted_table["cv_roc_auc_mean"],
            xerr=sorted_table["cv_roc_auc_std"], capsize=4,
            color=["#9ECAE9", "#6BAED6", "#3182BD", "#08519C"][-len(sorted_table):])
    ax.axvline(0.5, color="#E45756", linestyle="--", label="Chance reference (0.50)")
    ax.set_xlim(0, 1)
    ax.set(title="1. Model Selection: Stratified 5-Fold CV ROC-AUC",
           xlabel="Mean ROC-AUC on training folds", ylabel="")
    ax.legend()
    fig.tight_layout()
    p = PLOT_DIR / "01_cross_validation_model_comparison.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    # Confusion matrix with readable labels.
    cm = confusion_matrix(y_test, pred, labels=[0, 1])
    fig, ax = plt.subplots(figsize=(6.6, 5.4))
    sns.heatmap(cm, annot=True, fmt="d", cmap="Blues", cbar=False,
                xticklabels=[f"G3 < {PASS_MARK}", f"G3 >= {PASS_MARK}"],
                yticklabels=[f"G3 < {PASS_MARK}", f"G3 >= {PASS_MARK}"], ax=ax)
    ax.set(title="2. Held-Out Test Confusion Matrix", xlabel="Predicted class", ylabel="Actual class")
    fig.tight_layout()
    p = PLOT_DIR / "02_test_confusion_matrix.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    # ROC.
    fig, ax = plt.subplots(figsize=(7.2, 5.4))
    RocCurveDisplay.from_predictions(y_test, prob, name=f"{metrics['selected_model']} (AUC={metrics['test_roc_auc']:.3f})", ax=ax)
    ax.plot([0, 1], [0, 1], linestyle="--", color="gray", label="Chance")
    ax.set(title="3. ROC Curve on the Held-Out Test Set", xlabel="False-positive rate", ylabel="True-positive rate")
    ax.legend(loc="lower right")
    fig.tight_layout()
    p = PLOT_DIR / "03_test_roc_curve.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    # Metrics bars for selected model.
    metric_names = ["Accuracy", "Balanced accuracy", "Precision", "Recall", "F1", "ROC-AUC"]
    metric_values = [metrics["test_accuracy"], metrics["test_balanced_accuracy"], metrics["test_precision"],
                     metrics["test_recall"], metrics["test_f1"], metrics["test_roc_auc"]]
    fig, ax = plt.subplots(figsize=(9, 5.2))
    cols = ["#4C78A8", "#72B7B2", "#F58518", "#E45756", "#54A24B", "#B279A2"]
    bars = ax.bar(metric_names, metric_values, color=cols)
    ax.set_ylim(0, 1)
    ax.set_title(f"4. Held-Out Performance Metrics: {metrics['selected_model']}")
    ax.set_ylabel("Score (0–1)")
    ax.tick_params(axis="x", rotation=15)
    for bar, val in zip(bars, metric_values):
        ax.text(bar.get_x() + bar.get_width()/2, min(val + .025, .98), f"{val:.2f}", ha="center", fontsize=9)
    fig.tight_layout()
    p = PLOT_DIR / "04_test_performance_metrics.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    # Permutation importance on held-out data, using ROC AUC. This is a diagnostic,
    # not causal importance and may be unstable with a small sample.
    try:
        perm = permutation_importance(model, X_test, y_test, scoring="roc_auc",
                                      n_repeats=20, random_state=RANDOM_STATE, n_jobs=1)
        imp = pd.DataFrame({
            "feature": X_test.columns,
            "importance_mean": perm.importances_mean,
            "importance_std": perm.importances_std
        }).sort_values("importance_mean", ascending=False).head(12)
        imp.to_csv(OUT_DIR / "permutation_feature_importance.csv", index=False)
        display = imp.sort_values("importance_mean", ascending=True)
        fig, ax = plt.subplots(figsize=(9, 6.2))
        ax.barh(display["feature"], display["importance_mean"], xerr=display["importance_std"],
                color="#4C78A8", capsize=3)
        ax.axvline(0, color="gray", linewidth=1)
        ax.set(title="5. Permutation Importance on Test Data (ROC-AUC Decrease)",
               xlabel="Mean score decrease when feature is shuffled", ylabel="")
        fig.tight_layout()
        p = PLOT_DIR / "05_permutation_feature_importance.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        (OUT_DIR / "permutation_importance_warning.txt").write_text(str(exc), encoding="utf-8")
    return paths

def add_cell(cell, text, bold=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(9)
            run.bold = bold

def build_submission_description():
    text = SUBMISSION_TXT.read_text(encoding="utf-8").strip()
    doc = Document()
    sec = doc.sections[0]; sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.75)
    sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(11)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("WEEK 4 TASK"); run.bold = True; run.font.size = Pt(11); run.font.color.rgb = RGBColor(232, 106, 0)
    p = doc.add_heading("Submission Description", 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Machine Learning Model Development and Evaluation"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section in text.split("\n\n"):
        p = doc.add_paragraph(section)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.12
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"Word count: {len(text.split())}"); run.italic = True; run.font.size = Pt(9)
    doc.save(SUBMISSION_DOCX)

def build_report(df, audit, cv_table, metrics, images, source):
    doc = Document()
    sec = doc.sections[0]; sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.62)
    sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Title"].font.name = "Aptos Display"; doc.styles["Title"].font.size = Pt(23)
    doc.styles["Heading 1"].font.name = "Aptos Display"; doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.name = "Aptos Display"; doc.styles["Heading 2"].font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("WEEK 4 PROJECT REPORT"); run.bold = True; run.font.size = Pt(12); run.font.color.rgb = RGBColor(232, 106, 0)
    p = doc.add_heading("Machine Learning Model Development and Evaluation", 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("UCI Student Performance — Mathematics | Python and Scikit-learn"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Executive Summary", 1)
    doc.add_paragraph(
        f"This project trains classification models to estimate whether a student reaches the report's descriptive threshold of G3 >= {PASS_MARK}/20, "
        "using background and study-related variables while deliberately excluding G1 and G2 from the predictors to reduce target leakage. "
        f"The dataset contains {metrics['n_students_total']} eligible records; {metrics['positive_class_rate_all_pct']:.1f}% meet the threshold. "
        f"Four candidate models were compared by stratified five-fold cross-validation on the training split. "
        f"The selected model was {metrics['selected_model']}, with mean cross-validation ROC-AUC {metrics['selected_cv_roc_auc_mean']:.3f}. "
        f"On the held-out test set, accuracy was {metrics['test_accuracy']:.3f}, balanced accuracy {metrics['test_balanced_accuracy']:.3f}, "
        f"F1 {metrics['test_f1']:.3f}, and ROC-AUC {metrics['test_roc_auc']:.3f}. The test ROC-AUC is only modestly above chance, so the model has limited discrimination on this held-out sample. These metrics are educational evidence, not a validated high-stakes prediction system."
    )
    doc.add_heading("2. Problem Definition and Target", 1)
    doc.add_paragraph("Objective: demonstrate the end-to-end classification workflow, compare simple algorithms fairly, and discuss the strengths and limits of model evaluation.")
    doc.add_paragraph(f"Target label: 1 when final Mathematics grade G3 >= {PASS_MARK}; 0 when G3 < {PASS_MARK}. This binary threshold is an analytical convention for this assignment and is not a universal school policy.")
    doc.add_paragraph("Important leakage prevention: G1 and G2 are earlier grades closely related to G3, so they are excluded from all features. The goal is to test the value of non-grade predictors rather than letting earlier marks dominate the model. G3 is used only to construct the target and is never passed to the model.")
    doc.add_heading("3. Dataset and Preparation", 1)
    doc.add_paragraph(f"Dataset: UCI Machine Learning Repository, Student Performance Mathematics file. Page: {UCI_PAGE}. DOI: {DOI}. Acquisition: {source}.")
    doc.add_paragraph(f"Rows before preparation: {audit['source_rows']}; duplicate rows removed: {audit['duplicate_rows_removed']}; missing final grades: {audit['missing_final_grades_before_filter']}; invalid final grade values removed: {audit['invalid_final_grade_values_removed']}; records retained: {audit['rows_after_preparation']}.")
    doc.add_paragraph("Exact duplicate rows are removed, text fields are trimmed, known numeric columns are converted safely, and final grades are checked for the plausible 0–20 range. Rows without a target are excluded rather than assigning an invented outcome. Predictors are processed inside Scikit-learn pipelines: numeric missing values are median-imputed and scaled; categorical missing values are most-frequent-imputed and one-hot encoded. The preprocessor is fitted separately within each training fold.")
    doc.add_heading("4. Train/Test Design and Model Selection", 1)
    doc.add_paragraph(f"A stratified {int((1-TEST_SIZE)*100)}/{int(TEST_SIZE*100)} split was made with random_state={RANDOM_STATE}. The held-out test set contains {metrics['test_count']} records and is not used to choose a model. Stratified five-fold cross-validation was run on training data, with mean ROC-AUC specified as the selection metric.")
    doc.add_paragraph("Candidate algorithms:")
    for item in [
        "Dummy prior classifier: baseline reflecting the training-set class prevalence.",
        "Logistic Regression: a regularized linear baseline with class weighting.",
        "Decision Tree: a shallow, interpretable non-linear model with a minimum leaf size.",
        "Random Forest: an ensemble of trees, using class weighting and a minimum leaf size to limit excessive complexity."
    ]:
        doc.add_paragraph(item, style="List Bullet")
    table = doc.add_table(rows=1, cols=5); table.style = "Light Shading Accent 1"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, label in zip(table.rows[0].cells, ["Model", "Mean CV AUC", "CV AUC SD", "Mean CV F1", "CV balanced acc."]):
        add_cell(c, label, True)
    for _, row in cv_table.iterrows():
        cells = table.add_row().cells
        add_cell(cells[0], row["model"])
        add_cell(cells[1], f"{row['cv_roc_auc_mean']:.3f}")
        add_cell(cells[2], f"{row['cv_roc_auc_std']:.3f}")
        add_cell(cells[3], f"{row['cv_f1_mean']:.3f}")
        add_cell(cells[4], f"{row['cv_balanced_accuracy_mean']:.3f}")
    doc.add_heading("5. Held-Out Evaluation Metrics", 1)
    metric_table = doc.add_table(rows=1, cols=2); metric_table.style = "Light Shading Accent 1"
    add_cell(metric_table.rows[0].cells[0], "Metric", True); add_cell(metric_table.rows[0].cells[1], "Test score", True)
    for name, value in [
        ("Accuracy", metrics["test_accuracy"]),
        ("Balanced accuracy", metrics["test_balanced_accuracy"]),
        ("Precision for class 1", metrics["test_precision"]),
        ("Recall for class 1", metrics["test_recall"]),
        ("F1-score for class 1", metrics["test_f1"]),
        ("ROC-AUC", metrics["test_roc_auc"]),
    ]:
        cells = metric_table.add_row().cells
        add_cell(cells[0], name); add_cell(cells[1], f"{value:.3f}")
    doc.add_paragraph("Accuracy is the share of all test cases classified correctly. Precision describes the fraction of predicted-positive cases that are actual positives. Recall describes the fraction of actual-positive cases found. F1 combines precision and recall. Balanced accuracy averages sensitivity across the two classes; ROC-AUC measures ranking across classification thresholds. No metric alone is sufficient.")
    doc.add_page_break()
    doc.add_heading("6. Visualizations and Interpretation", 1)
    chart_notes = [
        ("Figure 1. Cross-validation model comparison", "Compares candidate models by training-only cross-validation ROC-AUC; error bars show fold-to-fold standard deviation. The highest mean score is used for model selection."),
        ("Figure 2. Test confusion matrix", "Shows true negatives, false positives, false negatives, and true positives on held-out test cases. False negatives may mean an at-risk student is missed; false positives may cause unnecessary follow-up."),
        ("Figure 3. Test ROC curve", "Shows the sensitivity-versus-false-positive-rate trade-off as the decision threshold changes. The diagonal is chance-level ranking."),
        ("Figure 4. Test performance metrics", "Presents complementary metrics on the same 0–1 scale; gaps between precision and recall can reveal different kinds of model error."),
        ("Figure 5. Permutation feature importance", "Measures how much test ROC-AUC changes when one input feature is shuffled. This is a model-dependence diagnostic, not a causal effect; correlated features and small samples can make rankings unstable.")
    ]
    for chart_index, (img, (title, note)) in enumerate(zip(images, chart_notes)):
        if chart_index > 0:
            doc.add_page_break()
        doc.add_heading(title, 2)
        doc.add_picture(str(img), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(note)
    doc.add_heading("7. Error Analysis, Overfitting, and Limitations", 1)
    cm = metrics["confusion_matrix"]
    doc.add_paragraph(f"The confusion matrix is ordered as actual/predicted classes [G3 < {PASS_MARK}, G3 >= {PASS_MARK}]. Its rows are {cm}. Inspect off-diagonal entries in the chart and consider whether false positives or false negatives are costlier for the intended context.")
    for item in [
        "Overfitting is limited by training-only cross-validation, a shallow decision tree, minimum leaf sizes, and a held-out test split. A large gap between training and validation performance would be a warning; this workflow emphasizes validation scores.",
        "The dataset is relatively small, and the final test metrics can vary with the random split. Cross-validation standard deviation is shown to make this variability visible.",
        "Class prevalence affects raw accuracy. Balanced accuracy, precision, recall, F1, and ROC-AUC provide additional perspectives.",
        "G1 and G2 are excluded to prevent the model from using earlier grades that might not be available at prediction time. If the practical use case includes those grades, a separate time-defined model could be evaluated, but it would answer a different question.",
        "Students come from two Portuguese secondary schools. Results may not generalize to other regions, schools, or later cohorts.",
        "Feature importance is not causal importance. Sensitive personal or family variables should not be used for real student decisions without domain review, fairness assessment, consent/privacy safeguards, and external validation."
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("8. Improvement Plan", 1)
    for item in [
        "Validate on a larger and more recent dataset from additional schools.",
        "Report confidence intervals for performance metrics and repeat the evaluation over several stratified splits.",
        "Check calibration and threshold choice in consultation with educators instead of assuming 0.5 is always optimal.",
        "Measure subgroup performance and error rates before considering any operational use.",
        "Compare against a simpler interpretable baseline and document the costs of false positives and false negatives."
    ]:
        doc.add_paragraph(item, style="List Number")
    doc.add_heading("9. Conclusion", 1)
    doc.add_paragraph(f"The selected model, {metrics['selected_model']}, was chosen using training-fold cross-validation and evaluated on a separate test set. The report records its metrics and errors transparently and uses five figures to explain model selection, predictive performance, and feature dependence. This is a reproducible learning exercise, not a tool for high-stakes individual decisions.")
    doc.add_heading("10. References", 1)
    doc.add_paragraph("Cortez, P. (2008). Student Performance [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5TG7T")
    doc.add_paragraph("Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.")
    doc.add_paragraph("Dataset page: " + UCI_PAGE)
    doc.save(REPORT_PATH)

def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    raw, source = acquire_data()
    df, audit = prepare_data(raw)
    excluded = {"G1", "G2", "G3", "target_pass"}
    X = df.drop(columns=[c for c in excluded if c in df.columns]).copy()
    y = df["target_pass"].astype(int)
    # Convert pandas string extension columns to object for compatibility with sklearn imputers.
    for col in X.select_dtypes(include=["string"]).columns:
        X[col] = X[col].astype(object)
    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=TEST_SIZE, random_state=RANDOM_STATE, stratify=y
    )
    model, pred, prob, cv_table, metrics = evaluate_models(X_train, y_train, X_test, y_test)
    metrics["n_students_total"] = int(len(df))
    metrics["positive_class_rate_all_pct"] = float(y.mean() * 100)
    metrics["source_features"] = list(X.columns)
    metrics["excluded_grade_features"] = ["G1", "G2", "G3"]
    metrics["model_selection_cv_auc_mean"] = float(cv_table.iloc[0]["cv_roc_auc_mean"])
    metrics["selected_cv_roc_auc_mean"] = float(cv_table.loc[cv_table["model"] == metrics["selected_model"], "cv_roc_auc_mean"].iloc[0])
    images = make_plots(y_test, pred, prob, cv_table, model, X_test, y_train, metrics)
    df.to_csv(OUT_DIR / "student_performance_modeling_data.csv", index=False)
    cv_table.to_csv(OUT_DIR / "cross_validation_model_comparison.csv", index=False)
    (OUT_DIR / "model_metrics.json").write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    (OUT_DIR / "data_quality_audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    (OUT_DIR / "run_metadata.json").write_text(json.dumps({
        "dataset": "UCI Student Performance — Mathematics",
        "source": source, "source_url": UCI_ZIP, "dataset_page": UCI_PAGE, "doi": DOI,
        "records_after_preparation": len(df), "synthetic_fallback_used": False,
        "target": f"1 if G3 >= {PASS_MARK}; otherwise 0",
        "excluded_features_to_reduce_leakage": ["G1", "G2", "G3"],
        "selected_model": metrics["selected_model"],
        "visualizations": [p.name for p in images],
        "report": REPORT_PATH.name,
    }, indent=2), encoding="utf-8")
    build_report(df, audit, cv_table, metrics, images, source)
    build_submission_description()
    print("Official UCI dataset confirmed; synthetic fallback not used.")
    print(f"Records: {len(df)}; positive class rate: {metrics['positive_class_rate_all_pct']:.1f}%")
    print("CV comparison:")
    print(cv_table.to_string(index=False))
    print(f"Selected model: {metrics['selected_model']}")
    print(f"Test accuracy={metrics['test_accuracy']:.3f}, F1={metrics['test_f1']:.3f}, ROC-AUC={metrics['test_roc_auc']:.3f}")
    print("Created report, submission DOCX, metrics JSON, audit data, and visualizations.")

if __name__ == "__main__":
    main()
