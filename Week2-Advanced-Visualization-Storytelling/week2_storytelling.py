"""Week 2: Advanced Data Visualization and Storytelling with Python.

Uses the official UCI Student Performance Mathematics CSV, preferably already
downloaded by the Week 1 project to ../data/student-mat.csv. If it is absent,
the script attempts to download the official UCI archive. It never substitutes
synthetic data for the official source.
"""
from io import BytesIO
from pathlib import Path
import json
import urllib.request
import zipfile

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
from scipy import stats
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
OUT_DIR = ROOT / "outputs"
PLOT_DIR = ROOT / "visualizations"
REPORT_PATH = ROOT / "Week2_Advanced_Data_Visualization_Report.docx"
SUBMISSION_DOCX = ROOT / "Submission_Description.docx"
SUBMISSION_TXT = ROOT / "submission_description.txt"

UCI_ZIP = "https://archive.ics.uci.edu/static/public/320/student+performance.zip"
UCI_PAGE = "https://archive.ics.uci.edu/dataset/320/student%2Bperformance"
DOI = "https://doi.org/10.24432/C5TG7T"
PASS_MARK = 10

def find_math_csv(zf):
    """Find student-mat.csv in a ZIP, including nested ZIP files."""
    for name in zf.namelist():
        if name.lower().endswith("student-mat.csv"):
            return zf.read(name)
    for name in zf.namelist():
        if name.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(BytesIO(zf.read(name))) as nested:
                    found = find_math_csv(nested)
                    if found is not None:
                        return found
            except zipfile.BadZipFile:
                continue
    return None

def acquire_data():
    """Load official UCI Mathematics data from Week 1 or acquire it directly."""
    candidates = [
        ROOT.parent / "data" / "student-mat.csv",
        ROOT / "data" / "student-mat.csv",
    ]
    for path in candidates:
        if path.exists():
            df = pd.read_csv(path, sep=";")
            if len(df) > 0 and {"G1", "G2", "G3", "studytime", "absences"}.issubset(df.columns):
                return df, "Existing official UCI Mathematics CSV from the Week 1 project"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(UCI_ZIP, headers={"User-Agent": "Week2-Student-EDA/1.0"})
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            archive_bytes = response.read()
        with zipfile.ZipFile(BytesIO(archive_bytes)) as archive:
            csv_bytes = find_math_csv(archive)
        if csv_bytes is None:
            raise RuntimeError("student-mat.csv not found in the downloaded UCI archive")
        (DATA_DIR / "student-mat.csv").write_bytes(csv_bytes)
        df = pd.read_csv(BytesIO(csv_bytes), sep=";")
        if len(df) == 0 or not {"G1", "G2", "G3"}.issubset(df.columns):
            raise RuntimeError("The downloaded file failed schema validation")
        return df, f"Downloaded official UCI archive: {UCI_ZIP}"
    except Exception as exc:
        raise RuntimeError("Could not acquire official UCI data; no synthetic data was used.") from exc

def clean_for_visualization(raw):
    """Light, documented validation without changing valid official records."""
    df = raw.copy()
    before = len(df)
    duplicates = int(df.duplicated().sum())
    df = df.drop_duplicates().copy()
    for col in df.select_dtypes(include=["object", "string"]).columns:
        df[col] = df[col].astype("string").str.strip()
    numeric = ["age", "Medu", "Fedu", "studytime", "failures", "absences", "G1", "G2", "G3"]
    for col in numeric:
        if col in df:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    required = ["G1", "G2", "G3", "studytime", "failures", "absences"]
    missing_required_before = int(df[required].isna().sum().sum())
    rows_before_required_drop = len(df)
    df = df.dropna(subset=required).copy()
    dropped_missing_rows = rows_before_required_drop - len(df)
    bounds = {
        "G1": (0, 20), "G2": (0, 20), "G3": (0, 20),
        "studytime": (1, 4), "failures": (0, 4), "absences": (0, 200),
    }
    invalid_mask = pd.Series(False, index=df.index)
    invalid_counts = {}
    for col, (lo, hi) in bounds.items():
        bad = ~df[col].between(lo, hi)
        invalid_counts[col] = int(bad.sum())
        invalid_mask |= bad
    invalid_rows = int(invalid_mask.sum())
    df = df.loc[~invalid_mask].copy()
    audit = {
        "rows_before": before,
        "duplicates_removed": duplicates,
        "missing_required_cells_before_drop": missing_required_before,
        "rows_dropped_for_missing_required_fields": dropped_missing_rows,
        "invalid_values_by_field": invalid_counts,
        "invalid_rows_removed": invalid_rows,
        "rows_after": len(df),
        "synthetic_fallback_used": False,
    }
    if df.empty:
        raise RuntimeError("No records remain after validation.")
    return df, audit

def calculate_metrics(df):
    grade_means = df[["G1", "G2", "G3"]].mean().to_dict()
    grade_medians = df[["G1", "G2", "G3"]].median().to_dict()
    study = df.groupby("studytime", observed=True)["G3"].agg(["mean", "median", "count", "std"])
    study.index = study.index.astype(int)
    failures = df.assign(pass_flag=df["G3"] >= PASS_MARK).groupby("failures", observed=True)["pass_flag"].agg(["mean", "count"])
    failures.index = failures.index.astype(int)
    return {
        "n_students": int(len(df)),
        "n_columns": int(len(df.columns)),
        "grade_mean": {k: float(v) for k, v in grade_means.items()},
        "grade_median": {k: float(v) for k, v in grade_medians.items()},
        "grade_std": {k: float(v) for k, v in df[["G1", "G2", "G3"]].std().to_dict().items()},
        "g3_min": float(df["G3"].min()),
        "g3_max": float(df["G3"].max()),
        "pass_mark": PASS_MARK,
        "pass_count": int((df["G3"] >= PASS_MARK).sum()),
        "below_pass_count": int((df["G3"] < PASS_MARK).sum()),
        "pass_rate_pct": float((df["G3"] >= PASS_MARK).mean() * 100),
        "g2_g3_corr": float(df["G2"].corr(df["G3"])),
        "absences_g3_corr": float(df["absences"].corr(df["G3"])),
        "failures_g3_corr": float(df["failures"].corr(df["G3"])),
        "studytime_groups": {
            str(int(i)): {
                "mean": float(row["mean"]), "median": float(row["median"]),
                "count": int(row["count"]), "std": None if pd.isna(row["std"]) else float(row["std"])
            } for i, row in study.iterrows()
        },
        "failure_groups": {
            str(int(i)): {"pass_rate_pct": float(row["mean"] * 100), "count": int(row["count"])}
            for i, row in failures.iterrows()
        },
    }

def add_caption(fig, text):
    fig.text(0.01, 0.005, text, ha="left", va="bottom", fontsize=8, wrap=True)
    fig.tight_layout(rect=[0, 0.07, 1, 1])

def make_visualizations(df, m):
    PLOT_DIR.mkdir(parents=True, exist_ok=True)
    sns.set_theme(style="whitegrid", context="notebook")
    paths = []

    fig, ax = plt.subplots(figsize=(9, 5.4))
    ax.hist(df["G3"], bins=np.arange(-0.5, 21.5, 1), color="#4C78A8", edgecolor="white")
    ax.axvline(m["grade_mean"]["G3"], color="#E45756", linewidth=2,
               label=f"Mean = {m['grade_mean']['G3']:.2f}")
    ax.axvline(PASS_MARK, color="#54A24B", linewidth=2, linestyle="--",
               label=f"Reference threshold = {PASS_MARK}/20")
    ax.set(title="1. Final Grade Distribution: Many Students Cluster Near the Threshold",
           xlabel="Final Mathematics grade (G3, 0–20)", ylabel="Number of students")
    ax.legend(frameon=True)
    add_caption(fig, "Caption: Final marks are shown on the full 0–20 scale; the reference line at 10/20 is used only as a descriptive threshold in this report.")
    p = PLOT_DIR / "01_final_grade_distribution.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    fig, ax = plt.subplots(figsize=(9, 5.4))
    long = df[["G1", "G2", "G3"]].melt(var_name="Period", value_name="Grade")
    sns.boxplot(data=long, x="Period", y="Grade", hue="Period", legend=False, ax=ax,
                palette=["#72B7B2", "#F2CF5B", "#4C78A8"])
    means = df[["G1", "G2", "G3"]].mean()
    for xpos, col in enumerate(["G1", "G2", "G3"]):
        ax.scatter(xpos, means[col], color="#E45756", marker="D", s=45, zorder=3)
        ax.text(xpos, means[col] + 0.55, f"Mean {means[col]:.1f}", ha="center", fontsize=8)
    ax.set(title="2. Grades Across the Three Assessment Periods",
           xlabel="Assessment period", ylabel="Grade (0–20)")
    add_caption(fig, "Caption: Box plots show medians, quartiles, and spread by period; red diamonds mark period means. This is a cohort-level comparison, not a matched improvement claim for every student.")
    p = PLOT_DIR / "02_grade_progression.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    fig, ax = plt.subplots(figsize=(9, 5.4))
    sns.regplot(data=df, x="G2", y="G3", ax=ax, scatter_kws={"alpha": 0.60, "s": 28},
                line_kws={"color": "#E45756", "linewidth": 2}, ci=95)
    ax.set(title=f"3. Second-Period Grade Is Strongly Associated with Final Grade (r = {m['g2_g3_corr']:.2f})",
           xlabel="Second-period grade (G2)", ylabel="Final grade (G3)")
    add_caption(fig, "Caption: Each point is one student. The fitted line and confidence band summarize a strong linear association; the chart is descriptive and is not evidence of a causal effect.")
    p = PLOT_DIR / "03_g2_vs_g3_relationship.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    study = df.groupby("studytime", observed=True)["G3"].agg(["mean", "count", "std"]).reindex([1, 2, 3, 4]).dropna()
    cis = []
    for _, row in study.iterrows():
        if row["count"] > 1 and pd.notna(row["std"]):
            margin = stats.t.ppf(0.975, int(row["count"] - 1)) * row["std"] / np.sqrt(row["count"])
        else:
            margin = 0.0
        cis.append(margin)
    fig, ax = plt.subplots(figsize=(9, 5.4))
    x = np.arange(len(study))
    ax.bar(x, study["mean"], yerr=cis, capsize=5, color=["#9ECAE9", "#6BAED6", "#3182BD", "#08519C"])
    ax.set_xticks(x, [f"{int(i)}\n(n={int(row['count'])})" for i, row in study.iterrows()])
    ax.set(title="4. More Study Time Is Associated with Higher Average Grades, but Not a Perfect Ladder",
           xlabel="Weekly study-time category (1 = under 2 hours; 4 = over 10 hours)", ylabel="Mean final grade (95% t interval)")
    add_caption(fig, "Caption: Bars show group means and error bars show 95% t confidence intervals. Group sizes vary, so read means together with n; self-reported study-time categories are not randomized treatments.")
    p = PLOT_DIR / "04_studytime_group_comparison.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    fig, ax = plt.subplots(figsize=(9, 5.4))
    hb = ax.hexbin(df["absences"], df["G3"], gridsize=22, mincnt=1, cmap="Blues")
    cb = fig.colorbar(hb, ax=ax); cb.set_label("Number of students")
    ax.set(title=f"5. Absences Alone Have Little Linear Relationship with G3 (r = {m['absences_g3_corr']:.2f})",
           xlabel="Recorded absences", ylabel="Final grade (G3, 0–20)")
    add_caption(fig, "Caption: Hexagons show where students are concentrated. A near-zero Pearson correlation does not rule out nonlinear patterns or confounding by other variables.")
    p = PLOT_DIR / "05_absences_grade_density.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    group = df.assign(pass_flag=(df["G3"] >= PASS_MARK).astype(int)).groupby("failures", observed=True)["pass_flag"].agg(["mean", "count"]).sort_index()
    fig, ax = plt.subplots(figsize=(9, 5.4))
    xs = np.arange(len(group))
    vals = group["mean"].to_numpy() * 100
    ax.bar(xs, vals, color="#F58518", edgecolor="white")
    ax.set_xticks(xs, [f"{int(i)} prior failure(s)\n(n={int(row['count'])})" for i, row in group.iterrows()])
    ax.set_ylim(0, 105)
    ax.set(title="6. Previous Course Failures Mark a Group with Greater Academic Risk",
           xlabel="Number of previous class failures", ylabel=f"Share with G3 ≥ {PASS_MARK}/20 (%)")
    for xval, val in zip(xs, vals):
        ax.text(xval, min(val + 2, 102), f"{val:.0f}%", ha="center", fontsize=9)
    add_caption(fig, "Caption: The chart compares the share reaching the report's 10/20 threshold by recorded prior failures. Small group sizes can make the most extreme categories less stable.")
    p = PLOT_DIR / "06_pass_rate_by_failures.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    cols = ["studytime", "failures", "absences", "G1", "G2", "G3"]
    corr = df[cols].corr()
    fig, ax = plt.subplots(figsize=(8.5, 6.5))
    sns.heatmap(corr, annot=True, fmt=".2f", cmap="vlag", vmin=-1, vmax=1, center=0,
                square=True, linewidths=.5, ax=ax)
    ax.set(title="7. Grade History Dominates the Linear Correlation Pattern")
    add_caption(fig, "Caption: Pearson coefficients summarize pairwise linear relationships. They do not control for confounders and should not be interpreted as causal effects.")
    p = PLOT_DIR / "07_correlation_heatmap.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    return paths

def build_submission_description():
    content = SUBMISSION_TXT.read_text(encoding="utf-8").strip()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.75)
    sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(11)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WEEK 2 TASK"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(232,106,0)
    p = doc.add_heading("Submission Description", 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Advanced Data Visualization and Storytelling with Python"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section in content.split("\n\n"):
        p = doc.add_paragraph(section); p.paragraph_format.space_after = Pt(10); p.paragraph_format.line_spacing = 1.12
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"Word count: {len(content.split())}"); r.italic = True; r.font.size = Pt(9)
    doc.save(SUBMISSION_DOCX)

def build_report(df, audit, metrics, paths, source):
    doc = Document()
    sec = doc.sections[0]; sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Title"].font.name = "Aptos Display"; doc.styles["Title"].font.size = Pt(24)
    doc.styles["Heading 1"].font.name = "Aptos Display"; doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.name = "Aptos Display"; doc.styles["Heading 2"].font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WEEK 2 PROJECT REPORT"); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(232,106,0)
    p = doc.add_heading("Advanced Data Visualization and Storytelling with Python", 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Student Performance — Mathematics | Python, Pandas, NumPy, Matplotlib, Seaborn")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Executive Summary", 1)
    doc.add_paragraph(
        f"This report tells a seven-chart story about {metrics['n_students']} students in the official UCI Mathematics dataset. "
        f"The mean final grade is {metrics['grade_mean']['G3']:.2f}/20, and {metrics['pass_count']} students "
        f"({metrics['pass_rate_pct']:.1f}%) meet the report's descriptive threshold of 10/20. "
        f"Prior-period grade G2 has a Pearson correlation of {metrics['g2_g3_corr']:.2f} with G3, while the linear "
        f"correlation between absences and G3 is {metrics['absences_g3_corr']:.2f}. "
        "These results provide context for targeted support but do not prove causal relationships."
    )
    doc.add_heading("2. Audience, Question, and Narrative", 1)
    doc.add_paragraph("Audience: educators, academic coordinators, and readers without a specialist statistics background.")
    doc.add_paragraph("Guiding question: What patterns in grades, study time, absences, and previous failures help describe academic outcomes, and which patterns deserve careful follow-up?")
    for item in [
        "Start with the final-grade distribution to establish the overall level and spread.",
        "Compare grade distributions across assessment periods to show the cohort-level progression.",
        "Highlight the relationship between second-period and final grades.",
        "Compare study-time groups while showing uncertainty and group sizes.",
        "Check absences through a density view rather than relying only on an average.",
        "Compare threshold attainment by previous failures, then summarize relationships in a correlation heatmap.",
    ]:
        doc.add_paragraph(item, style="List Number")
    doc.add_heading("3. Dataset and Method", 1)
    doc.add_paragraph(f"Source: UCI Machine Learning Repository, Student Performance. {UCI_PAGE}")
    doc.add_paragraph(f"DOI: {DOI}. Acquisition: {source}. The report uses the official Mathematics course CSV (student-mat.csv), not the Portuguese-language course file.")
    doc.add_paragraph(f"Initial source file: {audit['rows_before']} records. Records after duplicate removal and validation: {audit['rows_after']}. Duplicate rows removed: {audit['duplicates_removed']}; invalid rows excluded from charts: {audit['invalid_rows_removed']}. Synthetic data used: No.")
    doc.add_paragraph("The pipeline checks required grade/study/absence fields, converts numeric columns safely, removes exact duplicate rows, and validates plausible ranges before visualizing. It does not impute final grades because that would invent outcome data. Any records excluded for missing required values or invalid ranges are counted in the run metadata/audit.")
    doc.add_heading("4. Visualization Design and Findings", 1)
    doc.add_paragraph("Each figure is presented with its key finding and a rationale for the chosen visual form. The chart sequence moves from overall distribution to relationships, group comparisons, and a combined correlation summary.")
    captions = [
        ("Figure 1. Final-grade distribution", "The histogram shows the full grade range. The mean and 10/20 reference line give the reader two simple points of orientation before interpreting more detailed relationships.", f"Mean G3: {metrics['grade_mean']['G3']:.2f}; median: {metrics['grade_median']['G3']:.1f}; standard deviation: {metrics['grade_std']['G3']:.2f}. {metrics['pass_count']} students ({metrics['pass_rate_pct']:.1f}%) have G3 ≥ 10."),
        ("Figure 2. Grades across assessment periods", "Box plots are a compact way to compare medians, quartiles, variability, and potential extreme values. Red diamonds mark means, making central tendency easier to compare.", f"Mean G1: {metrics['grade_mean']['G1']:.2f}; mean G2: {metrics['grade_mean']['G2']:.2f}; mean G3: {metrics['grade_mean']['G3']:.2f}. Compare the distributions rather than assume every student's grade increased."),
        ("Figure 3. Second-period versus final grade", "A scatter plot preserves individual observations; a fitted line and confidence band expose the strength and direction of a linear association.", f"Pearson correlation between G2 and G3: {metrics['g2_g3_corr']:.2f}. G2 and G3 come from the same course timeline, so this relationship should not be described as causal."),
        ("Figure 4. Study-time comparison", "A bar chart with confidence intervals compares group means, and sample-size labels prevent a small group from looking as certain as a large one.", "Weekly study time is self-reported and categories are observational. Mean differences may reflect other variables, including prior attainment and support."),
        ("Figure 5. Absences and final-grade density", "Hexbin aggregates overlapping points into cells. This shows where students are concentrated and avoids a dense scatter plot becoming difficult to read.", f"Pearson correlation between absences and G3: {metrics['absences_g3_corr']:.2f}. A weak linear correlation does not rule out nonlinear relationships or hidden factors."),
        ("Figure 6. Pass-threshold rate by prior failures", "A bar chart converts the binary threshold outcome into percentages that are easy to compare. Sample sizes are included because failure groups may be unequal.", "Previous failures are an indicator of risk in these records, not a label for an individual student. Small groups can have unstable percentages."),
        ("Figure 7. Correlation heatmap", "A single heatmap summarizes the direction and magnitude of several pairwise linear relationships, making the strong grade-history pattern visible at a glance.", f"G2–G3 correlation: {metrics['g2_g3_corr']:.2f}; absences–G3 correlation: {metrics['absences_g3_corr']:.2f}; failures–G3 correlation: {metrics['failures_g3_corr']:.2f}. Correlation is not causation."),
    ]
    for figure_index, (path, (title, rationale, finding)) in enumerate(zip(paths, captions)):
        if figure_index > 0:
            doc.add_page_break()
        doc.add_heading(title, 2)
        doc.add_picture(str(path), width=Inches(5.9))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(); r = p.add_run("What the chart shows: "); r.bold = True; p.add_run(finding)
        p = doc.add_paragraph(); r = p.add_run("Why this visual: "); r.bold = True; p.add_run(rationale)
    doc.add_heading("5. Business / Scientific Implications", 1)
    doc.add_paragraph("For educators, the results support using multiple signals rather than relying on a single variable. Earlier-period grades may help identify students who could benefit from timely check-ins, and previous failure history can provide context for offering additional academic support. The weak overall linear relationship between absences and final marks in this sample suggests that absences alone would be a poor substitute for reviewing a student's broader circumstances.")
    doc.add_paragraph("Recommended follow-up is supportive, not punitive: combine teacher context, formative assessment, attendance records, and student input before deciding on an intervention. The charts describe associations in historical records and should not be used as a deterministic ranking of individual students.")
    doc.add_heading("6. Limitations and Responsible Interpretation", 1)
    for item in [
        "The dataset is observational; visual patterns do not prove causes.",
        "The records come from two Portuguese secondary schools, so results may not transfer directly to other education systems.",
        "Study time is reported in categories rather than measured as an exact number of hours.",
        "A threshold of G3 ≥ 10 is used as a convenient analytical indicator in this report, not as a universal institutional policy.",
        "G1 and G2 are earlier grades from the same course. They are informative for retrospective description but may create target leakage in a prediction model if their timing is not handled correctly.",
        "Small categories and outliers can affect the reliability of means and rates; sample sizes and uncertainty should be considered."
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("7. Reproducibility and Files", 1)
    doc.add_paragraph("Run python week2_storytelling.py after installing requirements.txt. The script writes seven figures under visualizations/, computed metrics and audit information under outputs/, this Word report, and a separate Word version of the submission description. The dataset is acquired from the official UCI archive or the checked Week 1 source file.")
    doc.add_heading("8. Conclusion", 1)
    doc.add_paragraph(f"The story is not that one factor explains grades; rather, grade history stands out in the correlation pattern, study-time groups show descriptive differences, and absences have a weak linear relationship with final grade in this cohort. Visuals, captions, group sizes, and limitations together make these findings clearer while discouraging unsupported causal claims.")
    doc.add_heading("9. References", 1)
    doc.add_paragraph("Cortez, P. (2008). Student Performance [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5TG7T")
    doc.add_paragraph("Cortez, P., & Silva, A. M. G. (2008). Using data mining to predict secondary school student performance. Proceedings of the 5th Annual Future Business Technology Conference.")
    doc.add_paragraph(f"Dataset page: {UCI_PAGE}")
    doc.save(REPORT_PATH)

def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    raw, source = acquire_data()
    df, audit = clean_for_visualization(raw)
    metrics = calculate_metrics(df)
    paths = make_visualizations(df, metrics)
    df.to_csv(OUT_DIR / "student_performance_for_visualization.csv", index=False)
    (OUT_DIR / "chart_metrics.json").write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    (OUT_DIR / "data_quality_audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    (OUT_DIR / "run_metadata.json").write_text(json.dumps({
        "dataset": "UCI Student Performance — Mathematics",
        "source": source, "source_url": UCI_ZIP, "dataset_page": UCI_PAGE, "doi": DOI,
        "n_records_analyzed": len(df), "synthetic_fallback_used": False,
        "report": REPORT_PATH.name, "visualizations": [p.name for p in paths]
    }, indent=2), encoding="utf-8")
    build_report(df, audit, metrics, paths, source)
    build_submission_description()
    print("Official UCI Student Performance data verified; synthetic fallback not used.")
    print(f"Records analyzed: {len(df)}")
    print(f"Mean G3: {metrics['grade_mean']['G3']:.2f}; threshold rate: {metrics['pass_rate_pct']:.1f}%")
    print(f"G2/G3 correlation: {metrics['g2_g3_corr']:.3f}; absences/G3 correlation: {metrics['absences_g3_corr']:.3f}")
    print(f"Generated {len(paths)} visualization files and both Word documents.")

if __name__ == "__main__":
    main()
