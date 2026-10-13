"""Week 3: Statistical Analysis and Hypothesis Testing in Python.

Primary hypothesis:
H0: Mean final grade among students with zero previous class failures is no
    greater than the mean among students with one or more previous failures.
H1: Mean final grade among students with zero previous class failures is higher.

Input is the official UCI Student Performance Mathematics CSV. No synthetic
fallback data is used. Results are descriptive evidence from observational data.
"""
from io import BytesIO
from pathlib import Path
import json
import math
import urllib.request
import zipfile

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
from scipy import stats
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
OUT_DIR = ROOT / "outputs"
PLOT_DIR = ROOT / "visualizations"
REPORT_PATH = ROOT / "Week3_Statistical_Analysis_Hypothesis_Testing_Report.docx"
SUBMISSION_DOCX = ROOT / "Submission_Description.docx"
SUBMISSION_TXT = ROOT / "submission_description.txt"

UCI_ZIP = "https://archive.ics.uci.edu/static/public/320/student+performance.zip"
UCI_PAGE = "https://archive.ics.uci.edu/dataset/320/student%2Bperformance"
DOI = "https://doi.org/10.24432/C5TG7T"
ALPHA = 0.05
PASS_MARK = 10

def find_math_csv(zf):
    for name in zf.namelist():
        if name.lower().endswith("student-mat.csv"):
            return zf.read(name)
    for name in zf.namelist():
        if name.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(BytesIO(zf.read(name))) as nested:
                    found = find_math_csv(nested)
                    if found is not None:
                        return found
            except zipfile.BadZipFile:
                continue
    return None

def acquire_data():
    candidates = [
        ROOT.parent / "data" / "student-mat.csv",
        ROOT / "data" / "student-mat.csv",
    ]
    for path in candidates:
        if path.exists():
            frame = pd.read_csv(path, sep=";")
            if len(frame) and {"G1", "G2", "G3", "failures", "studytime"}.issubset(frame.columns):
                return frame, "Existing official UCI Mathematics CSV from the Week 1 project"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(UCI_ZIP, headers={"User-Agent": "Week3-Student-Stats/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as response:
            archive_bytes = response.read()
        with zipfile.ZipFile(BytesIO(archive_bytes)) as archive:
            payload = find_math_csv(archive)
        if payload is None:
            raise RuntimeError("student-mat.csv not found in UCI archive")
        (DATA_DIR / "student-mat.csv").write_bytes(payload)
        frame = pd.read_csv(BytesIO(payload), sep=";")
        if frame.empty or not {"G1", "G2", "G3", "failures"}.issubset(frame.columns):
            raise RuntimeError("Downloaded data failed schema validation")
        return frame, "Downloaded official UCI archive: " + UCI_ZIP
    except Exception as exc:
        raise RuntimeError("Official UCI dataset could not be acquired; no synthetic data was substituted.") from exc

def prepare_data(raw):
    df = raw.copy()
    initial_rows = len(df)
    duplicates = int(df.duplicated().sum())
    df = df.drop_duplicates().copy()
    for col in df.select_dtypes(include=["object", "string"]).columns:
        df[col] = df[col].astype("string").str.strip()
    numeric = ["age", "studytime", "failures", "absences", "G1", "G2", "G3"]
    for col in numeric:
        if col in df:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    required = ["studytime", "failures", "absences", "G1", "G2", "G3"]
    missing_required_cells = int(df[required].isna().sum().sum())
    rows_before_na_drop = len(df)
    df = df.dropna(subset=required).copy()
    rows_removed_missing = rows_before_na_drop - len(df)
    ranges = {
        "studytime": (1, 4), "failures": (0, 4), "absences": (0, 200),
        "G1": (0, 20), "G2": (0, 20), "G3": (0, 20),
    }
    invalid_by_field = {}
    invalid_mask = pd.Series(False, index=df.index)
    for col, (lo, hi) in ranges.items():
        bad = ~df[col].between(lo, hi)
        invalid_by_field[col] = int(bad.sum())
        invalid_mask |= bad
    rows_removed_invalid = int(invalid_mask.sum())
    df = df.loc[~invalid_mask].copy()
    if df.empty:
        raise RuntimeError("No rows remain after cleaning and range validation.")
    df["failure_group"] = np.where(df["failures"].eq(0), "No previous failures", "One or more previous failures")
    df["passed_threshold"] = df["G3"].ge(PASS_MARK)
    audit = {
        "rows_before_cleaning": int(initial_rows),
        "duplicate_rows_removed": duplicates,
        "missing_cells_in_required_analysis_fields": missing_required_cells,
        "rows_removed_for_missing_required_fields": int(rows_removed_missing),
        "invalid_values_by_field": invalid_by_field,
        "rows_removed_for_invalid_ranges": rows_removed_invalid,
        "rows_after_cleaning": int(len(df)),
        "synthetic_fallback_used": False,
        "cleaning_note": "No outcome grades were imputed. Rows missing required test variables or with invalid ranges were excluded and counted."
    }
    return df, audit

def welch_df(a, b):
    a = np.asarray(a, dtype=float); b = np.asarray(b, dtype=float)
    va = np.var(a, ddof=1); vb = np.var(b, ddof=1)
    term_a = va / len(a); term_b = vb / len(b)
    numerator = (term_a + term_b) ** 2
    denominator = term_a ** 2 / (len(a) - 1) + term_b ** 2 / (len(b) - 1)
    return float(numerator / denominator)

def hedges_g(a, b):
    a = np.asarray(a, dtype=float); b = np.asarray(b, dtype=float)
    na, nb = len(a), len(b)
    pooled_var = (((na - 1) * np.var(a, ddof=1)) + ((nb - 1) * np.var(b, ddof=1))) / (na + nb - 2)
    pooled_sd = math.sqrt(pooled_var)
    d = (np.mean(a) - np.mean(b)) / pooled_sd if pooled_sd else float("nan")
    correction = 1 - (3 / (4 * (na + nb) - 9))
    return float(d * correction)

def analyze(df):
    nofail = df.loc[df["failures"].eq(0), "G3"].to_numpy(dtype=float)
    anyfail = df.loc[df["failures"].ge(1), "G3"].to_numpy(dtype=float)
    n0, n1 = len(nofail), len(anyfail)
    mean0, mean1 = float(nofail.mean()), float(anyfail.mean())
    sd0, sd1 = float(nofail.std(ddof=1)), float(anyfail.std(ddof=1))
    diff = mean0 - mean1
    df_welch = welch_df(nofail, anyfail)
    se_diff = math.sqrt(sd0 ** 2 / n0 + sd1 ** 2 / n1)
    t_stat = diff / se_diff
    p_welch = float(stats.ttest_ind(nofail, anyfail, equal_var=False, alternative="greater").pvalue)
    crit = float(stats.t.ppf(1 - ALPHA / 2, df_welch))
    ci_low, ci_high = diff - crit * se_diff, diff + crit * se_diff
    g = hedges_g(nofail, anyfail)
    mw = stats.mannwhitneyu(nofail, anyfail, alternative="greater", method="auto")
    ct = pd.crosstab(df["failure_group"], df["passed_threshold"]).reindex(
        index=["No previous failures", "One or more previous failures"], columns=[False, True], fill_value=0
    )
    chi2, p_chi, dof_chi, expected = stats.chi2_contingency(ct.to_numpy(), correction=False)
    n_chi = int(ct.to_numpy().sum())
    denom = min(ct.shape[0] - 1, ct.shape[1] - 1)
    cramers_v = math.sqrt(chi2 / (n_chi * denom)) if denom and n_chi else float("nan")
    study_groups = [gdf["G3"].to_numpy(dtype=float) for _, gdf in df.groupby("studytime", observed=True)]
    kw = stats.kruskal(*study_groups, nan_policy="omit")
    group_summary = {}
    for group_name, values in [
        ("No previous failures", nofail),
        ("One or more previous failures", anyfail),
    ]:
        group_summary[group_name] = {
            "n": int(len(values)), "mean": float(np.mean(values)),
            "sd": float(np.std(values, ddof=1)), "median": float(np.median(values)),
            "q1": float(np.quantile(values, 0.25)), "q3": float(np.quantile(values, 0.75)),
            "threshold_rate_pct": float(np.mean(values >= PASS_MARK) * 100),
        }
    studytime_summary = {}
    study_frame = df.groupby("studytime", observed=True)["G3"].agg(["count", "mean", "std", "median"])
    for cat, row in study_frame.iterrows():
        studytime_summary[str(int(cat))] = {
            "n": int(row["count"]), "mean": float(row["mean"]),
            "sd": float(row["std"]), "median": float(row["median"])
        }
    pass_rates = {}
    for gname, sub in df.groupby("failure_group", observed=True):
        passed = int(sub["passed_threshold"].sum())
        pass_rates[gname] = {"n": int(len(sub)), "passed": passed, "rate_pct": float(100 * passed / len(sub))}
    return {
        "n_students_analyzed": int(len(df)),
        "raw_final_grade_mean": float(df["G3"].mean()),
        "raw_final_grade_median": float(df["G3"].median()),
        "primary_hypothesis": {
            "h0": "mu_no_failures <= mu_one_or_more_failures",
            "h1": "mu_no_failures > mu_one_or_more_failures",
            "alpha": ALPHA,
            "test": "One-sided Welch independent-samples t-test",
            "no_failure_group": group_summary["No previous failures"],
            "prior_failure_group": group_summary["One or more previous failures"],
            "mean_difference_no_fail_minus_prior_fail": float(diff),
            "standard_error_difference": float(se_diff),
            "welch_t": float(t_stat),
            "welch_degrees_freedom": float(df_welch),
            "p_value_one_sided": p_welch,
            "ci95_difference_two_sided": [float(ci_low), float(ci_high)],
            "hedges_g": float(g),
            "interpretation_at_005": "Reject H0 in favor of H1" if p_welch < ALPHA else "Fail to reject H0",
        },
        "sensitivity_mann_whitney": {
            "u_statistic": float(mw.statistic), "p_value_one_sided": float(mw.pvalue),
            "interpretation_at_005": "Evidence supports higher grades in the no-failure group" if mw.pvalue < ALPHA else "Insufficient evidence at alpha 0.05"
        },
        "secondary_chi_square": {
            "table_rows": ["No previous failures", "One or more previous failures"],
            "table_columns": ["Below 10/20", "At or above 10/20"],
            "observed": ct.to_numpy().astype(int).tolist(),
            "expected": expected.tolist(),
            "chi_square": float(chi2), "degrees_freedom": int(dof_chi),
            "p_value": float(p_chi), "cramers_v": float(cramers_v),
            "interpretation_at_005": "Evidence of association" if p_chi < ALPHA else "Insufficient evidence of association"
        },
        "exploratory_kruskal_wallis_studytime": {
            "h_statistic": float(kw.statistic), "degrees_freedom": int(len(study_groups) - 1),
            "p_value": float(kw.pvalue),
            "interpretation_at_005": "At least one study-time grade distribution differs" if kw.pvalue < ALPHA else "Insufficient evidence of different grade distributions"
        },
        "failure_group_threshold_rates": pass_rates,
        "studytime_summary": studytime_summary,
        "assumption_notes": {
            "welch": "Independent student records are assumed; Welch's test does not require equal variances. Grade discreteness and skew are reviewed with distribution and Q-Q plots.",
            "chi_square": "Observed counts are shown; expected cell counts are checked in the JSON output. Independence of rows is assumed.",
            "multiple_testing": "The prior-failure Welch test is pre-specified primary; secondary tests are supportive or exploratory and are not treated as additional confirmatory hypotheses."
        }
    }

def add_caption(fig, caption):
    fig.text(0.01, 0.004, caption, ha="left", va="bottom", fontsize=8, wrap=True)
    fig.tight_layout(rect=[0, 0.075, 1, 1])

def make_visualizations(df, result):
    PLOT_DIR.mkdir(parents=True, exist_ok=True)
    sns.set_theme(style="whitegrid", context="notebook")
    paths = []
    primary = result["primary_hypothesis"]
    group_order = ["No previous failures", "One or more previous failures"]
    palette = {"No previous failures": "#4C78A8", "One or more previous failures": "#F58518"}

    fig, ax = plt.subplots(figsize=(9, 5.3))
    sns.violinplot(data=df, x="failure_group", y="G3", hue="failure_group", order=group_order,
                   palette=palette, inner="quartile", cut=0, legend=False, ax=ax)
    ax.set(title="1. Final Grade Distributions by Previous-Failure History",
           xlabel="", ylabel="Final Mathematics grade (G3, 0–20)")
    ax.tick_params(axis="x", rotation=8)
    add_caption(fig, "Caption: The violin shape summarizes the distribution; inner marks show quartiles. Group sizes and means are reported in the adjacent report text.")
    p = PLOT_DIR / "01_grade_distribution_by_failure_history.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    fig, ax = plt.subplots(figsize=(8.5, 5.3))
    means = [primary["no_failure_group"]["mean"], primary["prior_failure_group"]["mean"]]
    ns = [primary["no_failure_group"]["n"], primary["prior_failure_group"]["n"]]
    sds = [primary["no_failure_group"]["sd"], primary["prior_failure_group"]["sd"]]
    sems = [s / np.sqrt(n) for s, n in zip(sds, ns)]
    cis = [stats.t.ppf(0.975, n - 1) * sem for n, sem in zip(ns, sems)]
    x = np.arange(2)
    ax.bar(x, means, yerr=cis, capsize=6, color=[palette[g] for g in group_order], width=0.58)
    ax.set_xticks(x, [f"No previous failures\n(n={ns[0]})", f"One or more failures\n(n={ns[1]})"])
    ax.set_ylim(0, 20)
    ax.set(title="2. Mean Final Grades: Students with No Previous Failures Score Higher",
           ylabel="Mean final grade (95% t interval)", xlabel="")
    for xx, value in zip(x, means):
        ax.text(xx, value + 0.6, f"{value:.2f}", ha="center", fontweight="bold")
    add_caption(fig, f"Caption: Bars show means and 95% group-specific t intervals. The primary test uses Welch's unequal-variance standard error and a two-sided 95% interval for the mean difference of {primary['mean_difference_no_fail_minus_prior_fail']:.2f} points.")
    p = PLOT_DIR / "02_mean_grade_comparison_ci.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    fig, ax = plt.subplots(figsize=(8.5, 5.3))
    rates = result["failure_group_threshold_rates"]
    vals = [rates[g]["rate_pct"] for g in group_order]
    counts = [rates[g]["n"] for g in group_order]
    ax.bar(x, vals, color=[palette[g] for g in group_order], width=0.58)
    ax.set_xticks(x, [f"No previous failures\n(n={counts[0]})", f"One or more failures\n(n={counts[1]})"])
    ax.set_ylim(0, 100)
    ax.set(title=f"3. Share Reaching the Descriptive Threshold (G3 ≥ {PASS_MARK}/20)",
           ylabel="Students at or above threshold (%)", xlabel="")
    for xx, value in zip(x, vals):
        ax.text(xx, value + 2.0, f"{value:.1f}%", ha="center", fontweight="bold")
    add_caption(fig, "Caption: The threshold is an analytical indicator used for this project, not a universal school grading policy. The chi-square test assesses association between failure-history group and threshold attainment.")
    p = PLOT_DIR / "03_threshold_attainment_by_failure_history.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    fig, axes = plt.subplots(1, 2, figsize=(10, 4.8))
    for ax, gname, color in zip(axes, group_order, [palette[g] for g in group_order]):
        values = df.loc[df["failure_group"] == gname, "G3"].to_numpy(dtype=float)
        stats.probplot(values, dist="norm", plot=ax)
        ax.get_lines()[0].set_markerfacecolor(color)
        ax.get_lines()[0].set_markeredgecolor(color)
        ax.get_lines()[1].set_color("#E45756")
        ax.set_title("No previous failures" if gname == group_order[0] else "One or more previous failures")
        ax.set_xlabel("Theoretical normal quantiles")
        ax.set_ylabel("Ordered grade quantiles")
    add_caption(fig, "Caption: Q-Q plots compare group grade quantiles with a normal reference. Departures from the line signal non-normality; with these sample sizes, inference is also considered alongside the rank-based sensitivity test.")
    p = PLOT_DIR / "04_qq_diagnostics.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    study = df.groupby("studytime", observed=True)["G3"].agg(["mean", "count", "std"]).reindex([1, 2, 3, 4]).dropna()
    xs = np.arange(len(study))
    margins = []
    for _, row in study.iterrows():
        margin = stats.t.ppf(0.975, int(row["count"] - 1)) * row["std"] / np.sqrt(row["count"]) if row["count"] > 1 else 0
        margins.append(margin)
    fig, ax = plt.subplots(figsize=(9, 5.2))
    ax.bar(xs, study["mean"], yerr=margins, capsize=5, color=["#9ECAE9", "#6BAED6", "#3182BD", "#08519C"])
    ax.set_xticks(xs, [f"Category {int(i)}\n(n={int(row['count'])})" for i, row in study.iterrows()])
    ax.set_ylim(0, 20)
    ax.set(title="5. Final Grade by Weekly Study-Time Category",
           xlabel="1 = <2h; 2 = 2–5h; 3 = 5–10h; 4 = >10h per week",
           ylabel="Mean final grade (95% t interval)")
    add_caption(fig, "Caption: Study-time group means are accompanied by uncertainty intervals and sample sizes. The Kruskal-Wallis test compares full rank distributions; differences here are exploratory and do not prove study time causes marks to change.")
    p = PLOT_DIR / "05_grade_by_study_time.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)

    tests = [
        ("Welch t-test", primary["p_value_one_sided"]),
        ("Mann-Whitney U", result["sensitivity_mann_whitney"]["p_value_one_sided"]),
        ("Chi-square", result["secondary_chi_square"]["p_value"]),
        ("Kruskal-Wallis", result["exploratory_kruskal_wallis_studytime"]["p_value"]),
    ]
    fig, ax = plt.subplots(figsize=(9, 5.2))
    labels = [t[0] for t in tests]
    vals = [-np.log10(max(t[1], 1e-300)) for t in tests]
    colors = ["#4C78A8", "#72B7B2", "#F58518", "#B279A2"]
    ax.bar(np.arange(len(vals)), vals, color=colors)
    threshold = -np.log10(ALPHA)
    ax.axhline(threshold, color="#E45756", linestyle="--", label="p = 0.05 reference")
    ax.set_xticks(np.arange(len(labels)), labels, rotation=12, ha="right")
    ax.set_ylabel("−log10(p-value)")
    ax.set_title("6. Evidence Across Primary, Sensitivity, and Exploratory Tests")
    ax.legend()
    add_caption(fig, "Caption: Larger bars mean smaller p-values, not larger practical effects. Read this plot with effect sizes and confidence intervals; secondary tests are supportive/exploratory and not additional confirmatory hypotheses.")
    p = PLOT_DIR / "06_test_evidence_comparison.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    return paths

def add_cell_text(cell, value, bold=False):
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(9)
            run.bold = bold

def build_submission_description():
    content = SUBMISSION_TXT.read_text(encoding="utf-8").strip()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.75)
    sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(11)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WEEK 3 TASK"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(232, 106, 0)
    p = doc.add_heading("Submission Description", 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Statistical Analysis and Hypothesis Testing in Python"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for paragraph in content.split("\n\n"):
        p = doc.add_paragraph(paragraph)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.12
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"Word count: {len(content.split())}"); r.italic = True; r.font.size = Pt(9)
    doc.save(SUBMISSION_DOCX)

def build_report(df, audit, res, images, source):
    doc = Document()
    sec = doc.sections[0]; sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.62)
    sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Title"].font.name = "Aptos Display"; doc.styles["Title"].font.size = Pt(23)
    doc.styles["Heading 1"].font.name = "Aptos Display"; doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.name = "Aptos Display"; doc.styles["Heading 2"].font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WEEK 3 PROJECT REPORT"); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(232, 106, 0)
    p = doc.add_heading("Statistical Analysis and Hypothesis Testing in Python", 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("UCI Student Performance — Mathematics | SciPy, Pandas, NumPy, Matplotlib, Seaborn")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    ph = res["primary_hypothesis"]
    g0, g1 = ph["no_failure_group"], ph["prior_failure_group"]
    ci0, ci1 = ph["ci95_difference_two_sided"]
    doc.add_heading("1. Executive Summary", 1)
    doc.add_paragraph(
        f"This report tests whether students with no previous class failures have higher final Mathematics grades than students with one or more previous failures. "
        f"The analyzed sample contains {res['n_students_analyzed']} official UCI records. Mean G3 is {g0['mean']:.2f} for the no-failure group "
        f"(n={g0['n']}) and {g1['mean']:.2f} for the prior-failure group (n={g1['n']}). "
        f"The estimated mean difference is {ph['mean_difference_no_fail_minus_prior_fail']:.2f} grade points; Welch t={ph['welch_t']:.2f}, "
        f"df={ph['welch_degrees_freedom']:.1f}, one-sided p={ph['p_value_one_sided']:.3g}. "
        f"The two-sided 95% confidence interval for the difference is [{ci0:.2f}, {ci1:.2f}], and Hedges' g={ph['hedges_g']:.2f}. "
        f"At alpha={ALPHA:.2f}, the primary decision is: {ph['interpretation_at_005']}. Statistical evidence does not establish causation."
    )
    doc.add_heading("2. Research Question and Hypotheses", 1)
    doc.add_paragraph("Research question: In this dataset, do students with zero recorded previous class failures have higher mean final Mathematics grades than students with at least one recorded previous failure?")
    doc.add_paragraph("Group A: failures = 0. Group B: failures ≥ 1. Outcome: final-period Mathematics grade G3, recorded on a 0–20 scale.")
    doc.add_paragraph("Null hypothesis (H0): μA ≤ μB. Alternative hypothesis (H1): μA > μB. Significance level: α = 0.05, specified before interpreting the test.")
    doc.add_paragraph("This is a directional research hypothesis, so the primary Welch test uses a one-sided p-value. A two-sided 95% confidence interval is reported alongside it to quantify the estimated mean difference and its uncertainty.")
    doc.add_heading("3. Dataset, Data Quality, and Preparation", 1)
    doc.add_paragraph(f"Source: UCI Machine Learning Repository, Student Performance Mathematics file. Dataset page: {UCI_PAGE}. DOI: {DOI}. Acquisition method: {source}.")
    doc.add_paragraph(f"Rows before cleaning: {audit['rows_before_cleaning']}; exact duplicate rows removed: {audit['duplicate_rows_removed']}; missing cells in required analysis variables before filtering: {audit['missing_cells_in_required_analysis_fields']}; rows removed for missing required fields: {audit['rows_removed_for_missing_required_fields']}; rows removed for invalid ranges: {audit['rows_removed_for_invalid_ranges']}; rows analyzed: {audit['rows_after_cleaning']}. Synthetic data used: No.")
    doc.add_paragraph("The code trims text fields, converts the analysis variables to numeric types, removes exact duplicate records, and checks valid ranges for G1/G2/G3, study-time category, previous failures, and absences. It does not impute missing grade outcomes. All exclusions are recorded in the audit JSON.")
    doc.add_heading("4. Statistical Methodology", 1)
    doc.add_heading("4.1 Primary: Welch independent-samples t-test", 2)
    doc.add_paragraph("Welch's t-test compares the independent means of Group A and Group B without assuming equal population variances. Independence is assumed at the student-record level. Because grades are bounded and discrete, the report includes distribution plots, Q-Q diagnostics, and a rank-based sensitivity test.")
    doc.add_paragraph(f"The estimated difference is A − B = {ph['mean_difference_no_fail_minus_prior_fail']:.2f}. Standard error = {ph['standard_error_difference']:.2f}; t = {ph['welch_t']:.3f}; Welch degrees of freedom = {ph['welch_degrees_freedom']:.2f}; one-sided p = {ph['p_value_one_sided']:.5g}. Two-sided 95% CI for A − B = [{ci0:.2f}, {ci1:.2f}].")
    doc.add_heading("4.2 Effect size", 2)
    doc.add_paragraph(f"Hedges' g = {ph['hedges_g']:.3f}, a small-sample corrected standardized mean difference. The raw-point difference is more directly interpretable on the 0–20 grading scale; the standardized effect should be read together with the group spread.")
    doc.add_heading("4.3 Sensitivity: Mann–Whitney U", 2)
    mw = res["sensitivity_mann_whitney"]
    doc.add_paragraph(f"Mann–Whitney U tests whether values in Group A tend to be greater than values in Group B without relying on normally distributed outcomes. U = {mw['u_statistic']:.1f}; one-sided p = {mw['p_value_one_sided']:.5g}. This is a sensitivity analysis, not a test of means in general.")
    doc.add_heading("4.4 Secondary: chi-square test of independence", 2)
    chi = res["secondary_chi_square"]
    doc.add_paragraph(f"A threshold indicator (G3 ≥ {PASS_MARK}) is cross-tabulated with failure-history group. This threshold is an analysis convention, not a universal institutional grading rule. Pearson χ²({chi['degrees_freedom']}) = {chi['chi_square']:.3f}, p = {chi['p_value']:.5g}, Cramér's V = {chi['cramers_v']:.3f}. Expected cell counts are available in outputs/test_results.json and are checked before interpreting Pearson's chi-square approximation.")
    table = doc.add_table(rows=1, cols=3); table.style = "Light Shading Accent 1"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, txt in zip(table.rows[0].cells, ["Failure-history group", "Below 10/20", "At or above 10/20"]):
        add_cell_text(cell, txt, bold=True)
    for i, row_name in enumerate(chi["table_rows"]):
        cells = table.add_row().cells
        observed = chi["observed"][i]
        add_cell_text(cells[0], row_name)
        add_cell_text(cells[1], observed[0])
        add_cell_text(cells[2], observed[1])
    doc.add_heading("4.5 Exploratory: Kruskal–Wallis across study-time categories", 2)
    kw = res["exploratory_kruskal_wallis_studytime"]
    doc.add_paragraph(f"Kruskal–Wallis compares grade ranks across the four study-time categories: H({kw['degrees_freedom']}) = {kw['h_statistic']:.3f}, p = {kw['p_value']:.5g}. This result is exploratory; no unplanned pairwise post-hoc claims are made.")
    doc.add_heading("5. Visualizations and Interpretation", 1)
    captions = [
        ("Figure 1. Final grade distributions by previous-failure history", "The violin plot shows the shape and spread of grades in the two groups, complementing the mean-based primary test.", f"No-failure mean: {g0['mean']:.2f}; prior-failure mean: {g1['mean']:.2f}."),
        ("Figure 2. Mean final grades with confidence intervals", "Means and 95% group-specific t intervals make the central difference and group uncertainty visible. The inferential test uses Welch's unequal-variance method.", f"Estimated mean difference: {ph['mean_difference_no_fail_minus_prior_fail']:.2f}; 95% CI [{ci0:.2f}, {ci1:.2f}]."),
        ("Figure 3. Threshold attainment by failure history", "Percentages translate the final grade into a simple descriptive outcome for a non-technical audience. A chi-square test evaluates association between the two categorical variables.", f"No-failure threshold rate: {g0['threshold_rate_pct']:.1f}%; prior-failure rate: {g1['threshold_rate_pct']:.1f}%."),
        ("Figure 4. Q-Q diagnostics for grade groups", "Q-Q plots examine departures from a normal reference and make the rationale for adding a non-parametric sensitivity test explicit.", "The discrete 0–20 grade scale cannot be perfectly normal; tests are interpreted with group sizes, the estimated difference, confidence interval, and sensitivity analysis."),
        ("Figure 5. Final grade by weekly study-time category", "Group means with intervals and n labels show patterns without hiding unequal group sizes. The Kruskal–Wallis test compares the rank distributions across categories.", f"Kruskal–Wallis p = {kw['p_value']:.5g}; the result is exploratory and not evidence that more study time necessarily causes higher grades."),
        ("Figure 6. Evidence across statistical tests", "The negative log base 10 of p-value gives a compact view of statistical evidence; it does not measure practical importance. Effect size and confidence interval remain essential.", f"Primary Welch test p = {ph['p_value_one_sided']:.5g}; Hedges' g = {ph['hedges_g']:.2f}."),
    ]
    for idx, (img, (title, rationale, finding)) in enumerate(zip(images, captions)):
        if idx in (0, 2, 4):
            doc.add_page_break()
        doc.add_heading(title, 2)
        doc.add_picture(str(img), width=Inches(5.9))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(); rr = p.add_run("Result: "); rr.bold = True; p.add_run(finding)
        p = doc.add_paragraph(); rr = p.add_run("Why this visualization: "); rr.bold = True; p.add_run(rationale)
    doc.add_heading("6. Primary Result and Hypothesis Decision", 1)
    doc.add_paragraph(f"Decision at α = {ALPHA:.2f}: {ph['interpretation_at_005']}. The no-failure group's sample mean was {g0['mean']:.2f} compared with {g1['mean']:.2f} for students with at least one prior failure. The estimated difference was {ph['mean_difference_no_fail_minus_prior_fail']:.2f} points, with 95% CI [{ci0:.2f}, {ci1:.2f}]. Hedges' g was {ph['hedges_g']:.2f}.")
    if ph["p_value_one_sided"] < ALPHA:
        doc.add_paragraph("Because the pre-specified one-sided p-value is below 0.05, the null hypothesis is rejected in favor of the directional alternative for this dataset. This does not prove that previous failures cause lower final grades; prior failures may reflect earlier achievement, learning context, or other factors.")
    else:
        doc.add_paragraph("Because the pre-specified one-sided p-value is not below 0.05, the null hypothesis is not rejected. This means the analysis did not establish the directional claim at the selected significance level; it does not prove the groups are identical.")
    doc.add_paragraph(f"The sensitivity Mann–Whitney test p-value is {mw['p_value_one_sided']:.5g}; the secondary chi-square result is p = {chi['p_value']:.5g}. These are supplementary analyses and should not be conflated with the single primary confirmatory test.")
    doc.add_heading("7. Validity, Assumptions, and Limitations", 1)
    for item in [
        "Welch's t-test assumes independent observations; it does not require equal population variances. Q-Q plots and the Mann–Whitney sensitivity test help contextualize non-normal grade distributions.",
        "Students are observational units; class or school clustering may affect standard errors if records share unmodelled group-level factors.",
        "Previous failures are not randomly assigned. Confounding and reverse timing mean the group contrast is association, not a causal effect.",
        "A 10/20 cutoff is used only as a descriptive threshold and may not represent every institution's policy.",
        "The dataset represents students from two Portuguese secondary schools; findings may not generalize to other populations.",
        "Secondary chi-square and Kruskal–Wallis tests are exploratory/supportive. The p-value of the primary directional Welch test is pre-specified; effect sizes and confidence intervals are reported to avoid p-value-only conclusions."
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("8. Conclusion and Practical Implications", 1)
    doc.add_paragraph("The analysis provides a reproducible way to assess whether prior-failure history is associated with later academic outcomes. If the group difference is statistically significant, it can support reviewing how students who have experienced prior failure are offered timely, constructive academic support. It must not be used to label students or assume their future performance. Follow-up should account for prior grades, attendance, learning resources, and educator context, with suitable privacy protections.")
    doc.add_heading("9. Reproducibility and Deliverables", 1)
    doc.add_paragraph("Run python week3_hypothesis_testing.py after installing requirements.txt. The script creates this DOCX report, Submission_Description.docx, six PNG visualizations, cleaned analysis data, a data-quality audit, and a JSON file with test statistics, effect sizes, confidence intervals, contingency-table counts, and study-time summaries.")
    doc.add_heading("10. References", 1)
    doc.add_paragraph("Cortez, P. (2008). Student Performance [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5TG7T")
    doc.add_paragraph("Cortez, P., & Silva, A. M. G. (2008). Using data mining to predict secondary school student performance. Proceedings of the 5th Annual Future Business Technology Conference.")
    doc.add_paragraph("Dataset page: " + UCI_PAGE)
    doc.save(REPORT_PATH)

def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    raw, source = acquire_data()
    df, audit = prepare_data(raw)
    result = analyze(df)
    images = make_visualizations(df, result)
    df.to_csv(OUT_DIR / "student_performance_for_hypothesis_testing.csv", index=False)
    (OUT_DIR / "test_results.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    (OUT_DIR / "data_quality_audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    (OUT_DIR / "run_metadata.json").write_text(json.dumps({
        "dataset": "UCI Student Performance — Mathematics",
        "source": source, "source_url": UCI_ZIP, "dataset_page": UCI_PAGE,
        "doi": DOI, "records_analyzed": len(df),
        "synthetic_fallback_used": False, "report": REPORT_PATH.name,
        "visualizations": [p.name for p in images],
    }, indent=2), encoding="utf-8")
    build_report(df, audit, result, images, source)
    build_submission_description()
    p = result["primary_hypothesis"]
    print("Official UCI dataset confirmed; synthetic fallback not used.")
    print(f"Students analyzed: {len(df)}")
    print(f"No prior failures mean G3: {p['no_failure_group']['mean']:.3f} (n={p['no_failure_group']['n']})")
    print(f"One or more prior failures mean G3: {p['prior_failure_group']['mean']:.3f} (n={p['prior_failure_group']['n']})")
    print(f"Difference: {p['mean_difference_no_fail_minus_prior_fail']:.3f}; Welch t={p['welch_t']:.3f}; df={p['welch_degrees_freedom']:.2f}; one-sided p={p['p_value_one_sided']:.6g}")
    print(f"95% CI for difference: {p['ci95_difference_two_sided']}; Hedges g={p['hedges_g']:.3f}")
    print(f"Created {len(images)} figures, JSON statistics, the report, and the submission-description DOCX.")

if __name__ == "__main__":
    main()
