"""Week 1: acquire official UCI Student Performance data, clean, analyze and report.
No synthetic fallback is used: if the UCI archive cannot be reached and no
previously downloaded official Mathematics CSV exists, the script stops.
"""
from io import BytesIO
from pathlib import Path
import json, urllib.request, zipfile
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
OUT = ROOT / "outputs"
PLOTS = ROOT / "visualizations"
RAW = DATA / "student-mat.csv"
REPORT = ROOT / "Week1_Student_Performance_EDA_Report.docx"
ZIP_URL = "https://archive.ics.uci.edu/static/public/320/student+performance.zip"
DATA_URL = "https://archive.ics.uci.edu/dataset/320/student%2Bperformance"
DOI = "https://doi.org/10.24432/C5TG7T"
NUMERIC = ["age","Medu","Fedu","traveltime","studytime","failures","famrel","freetime","goout","Dalc","Walc","health","absences","G1","G2","G3"]
RANGES = {"age":(15,22),"Medu":(0,4),"Fedu":(0,4),"traveltime":(1,4),"studytime":(1,4),"failures":(0,4),"famrel":(1,5),"freetime":(1,5),"goout":(1,5),"Dalc":(1,5),"Walc":(1,5),"health":(1,5),"absences":(0,93),"G1":(0,20),"G2":(0,20),"G3":(0,20)}

def acquire():
    DATA.mkdir(parents=True, exist_ok=True)
    if RAW.exists():
        df = pd.read_csv(RAW, sep=";")
        if len(df) and "G3" in df:
            return df, "Official UCI student-mat.csv retained in data/"
        RAW.unlink()
    req = urllib.request.Request(ZIP_URL, headers={"User-Agent":"Student-Performance-EDA/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            archive_bytes = resp.read()
        def find_math_payload(zf):
            for entry in zf.namelist():
                if entry.lower().endswith("student-mat.csv"):
                    return zf.read(entry)
            # The UCI download can be a ZIP containing the student.zip archive.
            for entry in zf.namelist():
                if entry.lower().endswith(".zip"):
                    try:
                        with zipfile.ZipFile(BytesIO(zf.read(entry))) as nested:
                            payload = find_math_payload(nested)
                            if payload is not None:
                                return payload
                    except zipfile.BadZipFile:
                        continue
            return None
        with zipfile.ZipFile(BytesIO(archive_bytes)) as z:
            raw_bytes = find_math_payload(z)
        if raw_bytes is None:
            raise RuntimeError("student-mat.csv not found, including nested ZIP files")
        df = pd.read_csv(BytesIO(raw_bytes), sep=";")
        if df.empty or "G3" not in df:
            raise RuntimeError("Downloaded dataset failed validation")
        RAW.write_bytes(raw_bytes)
        return df, "Downloaded from official UCI ZIP archive: " + ZIP_URL
    except Exception as exc:
        raise RuntimeError("Official UCI acquisition failed; no demo data was substituted.") from exc

def clean(raw):
    d = raw.copy()
    before_rows, before_cols = d.shape
    for c in d.select_dtypes(include=["object","string"]).columns:
        d[c] = d[c].astype("string").str.strip()
    duplicates = int(d.duplicated().sum())
    d = d.drop_duplicates().copy()
    for c in NUMERIC:
        if c in d:
            d[c] = pd.to_numeric(d[c], errors="coerce")
    missing_before = int(d.isna().sum().sum())
    invalid = {}
    for c, (lo, hi) in RANGES.items():
        if c in d:
            mask = d[c].notna() & ~d[c].between(lo, hi)
            invalid[c] = int(mask.sum())
            d.loc[mask, c] = np.nan
    numeric_filled = categorical_filled = 0
    for c in d.select_dtypes(include=[np.number]).columns:
        n = int(d[c].isna().sum())
        if n:
            med = d[c].median()
            d[c] = d[c].fillna(0 if pd.isna(med) else med)
            numeric_filled += n
    for c in d.select_dtypes(include=["object","string","category"]).columns:
        n = int(d[c].isna().sum())
        if n:
            mode = d[c].mode(dropna=True)
            d[c] = d[c].fillna(mode.iloc[0] if len(mode) else "Unknown")
            categorical_filled += n
    audit = [("rows_before_cleaning",before_rows),("columns_before_cleaning",before_cols),("duplicate_rows_removed",duplicates),("rows_after_duplicate_removal",len(d)),("missing_cells_before_range_validation",missing_before),("invalid_values_marked_missing",sum(invalid.values())),("numeric_cells_imputed",numeric_filled),("categorical_cells_imputed",categorical_filled),("missing_cells_after_cleaning",int(d.isna().sum().sum()))]
    audit += [("invalid_values_"+k,v) for k,v in invalid.items()]
    details = {"rows_before":before_rows,"columns_before":before_cols,"duplicates":duplicates,"rows_after":len(d),"missing_before":missing_before,"invalid":sum(invalid.values()),"numeric_filled":numeric_filled,"categorical_filled":categorical_filled,"missing_after":int(d.isna().sum().sum())}
    return d, pd.DataFrame(audit,columns=["metric","value"]), details

def charts(d):
    PLOTS.mkdir(parents=True,exist_ok=True)
    sns.set_theme(style="whitegrid")
    made=[]
    fig,ax=plt.subplots(figsize=(8,5)); ax.hist(d.G3,bins=np.arange(-.5,21.5,1),edgecolor="black")
    ax.set(title="Distribution of Final Grades (G3)",xlabel="Final grade (0–20)",ylabel="Number of students")
    fig.tight_layout(); p=PLOTS/"01_final_grade_distribution.png"; fig.savefig(p,dpi=170); plt.close(fig); made.append(p)
    g=d.groupby("studytime")["G3"].agg(["mean","count"]).reindex([1,2,3,4]).dropna()
    fig,ax=plt.subplots(figsize=(8,5)); bars=ax.bar(g.index.astype(str),g["mean"])
    ax.set(title="Average Final Grade by Weekly Study Time",xlabel="Study-time category (1 = <2h, 4 = >10h/week)",ylabel="Average grade (0–20)")
    for b,n in zip(bars,g["count"]): ax.annotate("n="+str(int(n)),(b.get_x()+b.get_width()/2,b.get_height()),xytext=(0,4),textcoords="offset points",ha="center",fontsize=8)
    fig.tight_layout(); p=PLOTS/"02_grade_by_studytime.png"; fig.savefig(p,dpi=170); plt.close(fig); made.append(p)
    fig,ax=plt.subplots(figsize=(8,5)); sns.scatterplot(data=d,x="absences",y="G3",alpha=.68,ax=ax)
    ax.set(title="School Absences vs Final Grade",xlabel="Number of school absences",ylabel="Final grade (0–20)")
    fig.tight_layout(); p=PLOTS/"03_absences_vs_grade.png"; fig.savefig(p,dpi=170); plt.close(fig); made.append(p)
    cols=[c for c in ["studytime","failures","absences","G1","G2","G3"] if c in d]
    fig,ax=plt.subplots(figsize=(8,6)); sns.heatmap(d[cols].corr(),annot=True,fmt=".2f",vmin=-1,vmax=1,center=0,cmap="coolwarm",ax=ax)
    ax.set_title("Correlation Among Study Habits, Absences and Grades")
    fig.tight_layout(); p=PLOTS/"04_numeric_correlation_heatmap.png"; fig.savefig(p,dpi=170); plt.close(fig); made.append(p)
    return made

def metrics(d):
    g=d.G3
    grouped=d.groupby("studytime")["G3"].agg(["mean","count"]).dropna()
    best=int(grouped["mean"].idxmax())
    return {"n_students":int(len(d)),"n_columns":int(len(d.columns)),"g3_mean":float(g.mean()),"g3_median":float(g.median()),"g3_std":float(g.std()),"g3_min":float(g.min()),"g3_max":float(g.max()),"pass_count":int((g>=10).sum()),"pass_rate_pct":float((g>=10).mean()*100),"below_count":int((g<10).sum()),"absences_g3_corr":float(d.absences.corr(g)),"g1_g3_corr":float(d.G1.corr(g)),"g2_g3_corr":float(d.G2.corr(g)),"best_studytime":best,"best_studytime_mean":float(grouped.loc[best,"mean"]),"studytime_groups":{str(int(i)):{"mean":float(r["mean"]),"count":int(r["count"])} for i,r in grouped.iterrows()}}

def report(d, audit, details, m, plots, source):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(10)
    doc.styles["Title"].font.name="Aptos Display"; doc.styles["Title"].font.size=Pt(24)
    doc.styles["Heading 1"].font.name="Aptos Display"; doc.styles["Heading 1"].font.size=Pt(16)
    doc.styles["Heading 2"].font.name="Aptos Display"; doc.styles["Heading 2"].font.size=Pt(12)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("WEEK 1 PROJECT REPORT"); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(232,106,0)
    p=doc.add_heading("Data Acquisition, Cleaning, and Exploratory Data Analysis",0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph("UCI Student Performance — Mathematics course | Python, Pandas, NumPy, Matplotlib, Seaborn"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("1. Objective",1)
    doc.add_paragraph("This project demonstrates data acquisition, data-quality inspection, cleaning, summary statistics, visualization, and interpretation using a publicly available student-performance dataset. The analysis focuses on final grade (G3) and its relationship with study time, school absences, and previous period grades.")
    doc.add_heading("2. Dataset and acquisition",1)
    doc.add_paragraph("The Mathematics file student-mat.csv is part of the UCI Machine Learning Repository Student Performance dataset. It contains grade records, study habits, absences, demographics, and family/social context for Portuguese secondary-school students. The script downloads the UCI ZIP archive and extracts the official Mathematics CSV; synthetic fallback data is not used.")
    doc.add_paragraph("Official source: "+DATA_URL+"\nDOI: "+DOI+"\nAcquisition: "+source+"\nInitial table size: "+str(details["rows_before"])+" rows × "+str(details["columns_before"])+" columns.")
    doc.add_heading("3. Methodology and cleaning decisions",1)
    for item in ["Acquire the source archive and load the Mathematics course CSV using its semicolon delimiter.","Inspect dimensions, column types, duplicate rows, and missing cells.","Trim whitespace in text fields and safely convert numeric fields.","Check plausible ranges for grades, absences, study-time ratings, age, and scaled questionnaire variables; invalid values are treated as missing.","Remove exact duplicate rows; impute numeric missing values with column medians and categorical missing values with the most frequent category (or Unknown if no mode exists).","Export the cleaning audit and summary statistics for reproducibility."]:
        doc.add_paragraph(item,style="List Number")
    doc.add_paragraph("These checks are retained even though the UCI dataset documentation indicates that the original data has no missing values. No valid record is dropped solely because a value is missing; the script records every duplicate, invalid value, and imputation in outputs/cleaning_audit.csv.")
    doc.add_heading("Cleaning audit",2)
    table=doc.add_table(rows=1,cols=2); table.style="Light Shading Accent 1"; table.rows[0].cells[0].text="Metric"; table.rows[0].cells[1].text="Value"
    for _,row in audit.iterrows():
        cells=table.add_row().cells; cells[0].text=str(row["metric"]); cells[1].text=str(row["value"])
    doc.add_heading("4. Exploratory data analysis",1)
    doc.add_paragraph("G3 is the final Mathematics grade on a 0–20 scale. For descriptive reporting only, G3 ≥ 10 is used as an operational pass threshold; this should not be taken as a statement about every institution's grading policy.")
    descriptions=[
      ("Figure 1. Final-grade distribution","The histogram shows how final marks are spread across the 0–20 range."),
      ("Figure 2. Average final grade by weekly study time","Group means compare the four study-time categories; sample sizes are printed above bars. Differences are associations, not proof of causation."),
      ("Figure 3. Absences versus final grade","The scatter plot shows the distribution of marks by recorded absences and helps reveal variation and potential outliers."),
      ("Figure 4. Numeric correlation heatmap","The heatmap compares linear correlations among study time, failures, absences, and grades. Correlation does not establish causation.")
    ]
    for p,(title,desc) in zip(plots,descriptions):
        doc.add_heading(title,2); doc.add_picture(str(p),width=Inches(5.8)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph(desc)
    doc.add_heading("5. Results and insights from the official data",1)
    doc.add_paragraph(f"After cleaning, the dataset contains {m['n_students']} students and {m['n_columns']} columns. G3 has a mean of {m['g3_mean']:.2f}, a median of {m['g3_median']:.2f}, a standard deviation of {m['g3_std']:.2f}, and a range of {m['g3_min']:.0f}–{m['g3_max']:.0f}.")
    doc.add_paragraph(f"With the report's G3 ≥ 10 pass indicator, {m['pass_count']} students ({m['pass_rate_pct']:.1f}%) meet the threshold and {m['below_count']} are below it.")
    doc.add_paragraph(f"Study-time category {m['best_studytime']} has the highest average final grade among represented categories: {m['best_studytime_mean']:.2f}. Categories mean: 1 = less than 2 hours, 2 = 2–5 hours, 3 = 5–10 hours, and 4 = more than 10 hours per week.")
    doc.add_paragraph(f"Pearson correlations with G3 are {m['absences_g3_corr']:.3f} for absences, {m['g1_g3_corr']:.3f} for G1, and {m['g2_g3_corr']:.3f} for G2. G1 and G2 are earlier grades from the same course, so their correlation with final marks is expected and creates a target-leakage risk if used for predictions made before those grades exist.")
    doc.add_heading("6. Limitations",1)
    for item in ["The records are observational; correlations cannot prove causal effects.","Study-time groups can differ in other factors, and group sizes should be considered.","G1 and G2 are closely linked to G3 by the grading timeline and should be avoided in prediction models intended to operate before those marks are available.","Data from two Portuguese secondary schools should not automatically be generalized to all students or countries."]:
        doc.add_paragraph(item,style="List Number")
    doc.add_heading("7. Reproducibility and outputs",1)
    doc.add_paragraph("Install packages from requirements.txt and run python student_performance_eda.py. The script saves the official source CSV in data/, cleaned data and audit/statistics files in outputs/, four PNG visualizations in visualizations/, metadata proving that synthetic fallback was not used, and this DOCX report in the project root. If UCI retrieval fails and no previously downloaded official CSV exists, the script exits with an error rather than generating substitute findings.")
    doc.add_heading("8. Conclusion",1)
    doc.add_paragraph("The project demonstrates a reproducible workflow for acquiring, cleaning, and exploring a real public dataset. The cleaning audit makes transformations transparent, while grade-distribution, study-time, absence, and correlation visualizations provide complementary views of student achievement. The findings are descriptive and should be interpreted with due attention to observational-data limitations.")
    doc.add_heading("9. References",1)
    doc.add_paragraph("Cortez, P. (2008). Student Performance [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5TG7T")
    doc.add_paragraph("Cortez, P., & Silva, A. M. G. (2008). Using data mining to predict secondary school student performance. Proceedings of the 5th Annual Future Business Technology Conference.")
    doc.add_paragraph("Dataset page: "+DATA_URL)
    doc.save(REPORT)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    raw,source=acquire()
    d,audit,details=clean(raw)
    m=metrics(d)
    d.to_csv(OUT/"student_performance_cleaned.csv",index=False)
    audit.to_csv(OUT/"cleaning_audit.csv",index=False)
    d.describe(include="all").transpose().to_csv(OUT/"summary_statistics.csv")
    (OUT/"eda_metrics.json").write_text(json.dumps(m,indent=2),encoding="utf-8")
    (OUT/"run_metadata.json").write_text(json.dumps({"dataset":"UCI Student Performance — Mathematics","source":source,"source_url":ZIP_URL,"dataset_page":DATA_URL,"doi":DOI,"raw_file":"data/student-mat.csv","raw_rows":int(details["rows_before"]),"clean_rows":int(details["rows_after"]),"synthetic_fallback_used":False},indent=2),encoding="utf-8")
    ps=charts(d); report(d,audit,details,m,ps,source)
    print("Dataset: official UCI Student Performance — Mathematics")
    print("Acquisition:",source)
    print(f"Rows: {details['rows_before']} -> {details['rows_after']}; missing after cleaning: {details['missing_after']}")
    print(f"Mean G3: {m['g3_mean']:.2f}; median G3: {m['g3_median']:.2f}; pass indicator rate: {m['pass_rate_pct']:.1f}%")
    print(f"Absences/G3 correlation: {m['absences_g3_corr']:.3f}")
    print("Created:",REPORT)

if __name__=="__main__":
    main()
